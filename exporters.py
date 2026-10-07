"""スライド(pptx)・PDF出力."""
import io
import os
import tempfile

FONT_JP = "UD Digi Kyokasho N-R"
FONT_JP_BOLD = "UD Digi Kyokasho N-B"
FONT_TTC_R = r"C:\Windows\Fonts\UDDigiKyokashoN-R.ttc"
FONT_TTC_B = r"C:\Windows\Fonts\UDDigiKyokashoN-B.ttc"


def build_pptx(title, problems, test_name=""):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    # タイトルスライド
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    tx = slide.shapes.title if slide.shapes.title else slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(12), Inches(1)).text_frame
    tx.text = title
    # 問題スライド (1問1スライド、最大10枚)
    for pb in problems[:10]:
        s = prs.slides.add_slide(prs.slide_layouts[5])
        tb = s.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12.3), Inches(4)).text_frame
        tb.word_wrap = True
        p = tb.paragraphs[0]
        p.text = f"問{pb['no']} [{pb.get('kind','')}]"
        for line in pb["q"].split("\n")[:6]:
            p2 = tb.add_paragraph()
            p2.text = line
        # ヒント
        ph = s.shapes.add_textbox(Inches(0.5), Inches(4.6), Inches(12.3), Inches(1)).text_frame
        ph.word_wrap = True
        ph.paragraphs[0].text = "ヒント: " + pb.get("hint", "")
        # 解答 (発表者ノート相当として最終テキスト)
        pa = s.shapes.add_textbox(Inches(0.5), Inches(5.8), Inches(12.3), Inches(1)).text_frame
        pa.word_wrap = True
        pa.paragraphs[0].text = "答: " + pb.get("answer", "")
    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)
    return buf


def _register_fonts():
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    try:
        pdfmetrics.registerFont(TTFont("Kyokasho", FONT_TTC_R, subfontIndex=0))
    except Exception:
        pass
    try:
        pdfmetrics.registerFont(TTFont("Kyokasho-B", FONT_TTC_B, subfontIndex=0))
    except Exception:
        pass


def build_pdf(title, problems, meta_line=""):
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib import colors
    from reportlab.lib.units import mm
    _register_fonts()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=15*mm, bottomMargin=15*mm)
    from reportlab.pdfbase.pdfmetrics import getRegisteredFontNames
    font = "Kyokasho" if "Kyokasho" in getRegisteredFontNames() else "Helvetica"
    st_title = ParagraphStyle("t", fontName=font, fontSize=16, leading=22)
    st_h = ParagraphStyle("h", fontName=font, fontSize=11, leading=16)
    st_b = ParagraphStyle("b", fontName=font, fontSize=10, leading=15)
    story = [Paragraph(title, st_title), Paragraph(meta_line, st_h), Spacer(1, 6*mm)]
    for pb in problems:
        q = pb["q"].replace("\n", "<br/>")
        story.append(Paragraph(f"問{pb['no']} [{pb.get('kind','')}]<br/>{q}", st_b))
        story.append(Spacer(1, 2*mm))
        story.append(Paragraph(f"ヒント: {pb.get('hint','')}<br/>答: {pb.get('answer','')}", st_b))
        story.append(Spacer(1, 4*mm))
    # ルーブリック
    story.append(Paragraph("【観点別ルーブリック】知識・技能60 / 思考・判断・表現30 / 主体性10", st_h))
    doc.build(story)
    buf.seek(0)
    return buf
